import math
import numpy as np
from pptx import Presentation
from pptx.util import Inches
import deck as D
from scene import Scene, render_pptx, render_preview
from deck import Fig, Ax, Se, Sde, TT, TC, GRAY, BLUE, ORANGE, TEAL, RED, INK, NAVY, BG, MUTED, FAINT

D.S = Scene(); S = D.S
PLUM = "#7A3E9D"
EQ = "#5B6B86"

class EqAx(Ax):
    """Ejes con la misma escala (pulgadas/unidad) en x e y: necesario para dibujar giros de solido rigido."""
    def __init__(self, fig, box, xlim, s, ybot):
        l, b, w, h = box
        yspan = (h * fig.H) / ((w * fig.W) / (xlim[1] - xlim[0]))
        super().__init__(fig, box, xlim, (ybot, ybot + yspan))

def pier_pts(H=2.6, B0=1.0, B1=0.7):
    o = (B0 - B1) / 2
    return [(0, 0), (B0, 0), (B0 - o, H), (o, H)]
def rot(p, th, piv=(1.0, 0.0)):
    dx, dz = p[0] - piv[0], p[1] - piv[1]
    return (piv[0] + dx * math.cos(th) + dz * math.sin(th), piv[1] - dx * math.sin(th) + dz * math.cos(th))

def draw_pier(A, th, loads=True, color=INK):
    g = pier_pts()
    gl = g + [g[0]]
    A.poly([p[0] for p in gl], [p[1] for p in gl], FAINT, 0.8, clip=False)
    A.dot((0.5, 2.6), 0.035, INK, hollow=True, bg=BG)
    r = [rot(p, th) for p in g]; rl = r + [r[0]]
    A.poly([p[0] for p in rl], [p[1] for p in rl], color, 2.4, clip=False)
    A.line((-0.3, 0), (1.5, 0), INK, 1.5)
    A.dot((1.0, 0.0), 0.03, INK)
    if loads:
        for t in (0.2, 0.45, 0.7, 0.92):
            p = rot((1.0 - 0.15 * t, 2.6 * t), th)      # borde derecho
            A.line((p[0] + 0.05, p[1]), (p[0] + 0.5, p[1]), RED, 1.5, arrow="end")
    c = rot((0.5, 2.6), th); A.dot(c, 0.042, BLUE)
    return c

def lam_curve(B, mark=True):
    B.spines("d_C", "λ")
    B.poly([0, 1], [0.8, 0.0], GRAY, 2.4)
    B.dot((0, 0.8), 0.03, INK)
    B.line((0, 0.8), (-0.0, 0.8), INK, 1)
    B.text((0, 0.8), "λ_0", dx=-0.05, ha="right", size=9.5, italic=True, box_w=0.6)
    B.line((1, 0), (1, 0.06), INK, 1); B.dot((1, 0), 0.03, INK)
    B.text((1, 0), "d_{C0}", dy=0.13, ha="center", size=9.5, italic=True, box_w=0.8)

def panel1(f):
    A = EqAx(f, [0.0, 0.2, 0.44, 0.78], (-0.8, 3.2), 0.42, -1.0)
    th = 0.2; c = draw_pier(A, th)
    top = 2.6 + 0.3
    A.line((0.5, 2.6), (0.5, top + 0.1), INK, 0.8, "dot"); A.line((c[0], c[1]), (c[0], top + 0.1), INK, 0.8, "dot")
    A.line((0.5, top), (c[0], top), BLUE, 1.3, arrow="both")
    A.text(((0.5 + c[0]) / 2, top), "d_C", color=BLUE, bold=True, italic=True, size=11, ha="center", va="bottom", box_w=0.8, dy=-0.02)
    A.line((1.9, -0.45), (-0.1, -0.45), RED, 3.4, arrow="end", big=True)
    A.text((0.9, -0.82), "F_b = λ W", color=RED, bold=True, italic=True, size=11, ha="center", box_w=1.3)
    B = Ax(f, [0.6, 0.34, 0.38, 0.56], (0, 1.3), (0, 1.25)); lam_curve(B)
    f.text(0.5, 0.05, "F_i ∝ W_i   (aceleración uniforme)", ha="center", size=9, box_w=3.6)

def panel2(f):
    A = Ax(f, [0.14, 0.36, 0.82, 0.58], (0, 1.2), (0, 1.25)); A.spines("d_C, d*", "λ, a*")
    G, k = 1.35, 1.3
    A.poly([0, 1], [0.8, 0.0], GRAY, 2.0)
    A.poly([0, 1 / G], [0.8 * k, 0.0], BLUE, 2.6)
    A.dot((0, 0.8 * k), 0.03, BLUE); A.dot((1 / G, 0), 0.03, BLUE)
    A.line((0, 0.8 * k), (0.0, 0.8 * k), BLUE, 1)
    A.text((0, 0.8 * k), "a_0^*", color=BLUE, dx=-0.05, ha="right", size=9.5, italic=True, box_w=0.7)
    A.text((1 / G, 0), "d_0^*", color=BLUE, dy=0.13, ha="center", size=9.5, italic=True, box_w=0.7)
    f.text(0.52, 0.955, "1GDL* = ESTRUCTURA/Γ", color="#0070C0", bold=True, italic=True, size=9, box_w=1.6)
    S.line(f.pt(0.62, 0.9), f.pt(0.45, 0.74), INK, 1, dash="dash", arrow="end")
    f.text(0.5, 0.105, "d* = d_C/Γ     a* = λ g / e*     m* = Σ W_i Φ_i / g", ha="center", size=9, italic=True, box_w=3.8)
    f.text(0.5, 0.03, "Γ = Σ W_i Φ_i / Σ W_i Φ_i^2     e* = Γ m* / M", ha="center", size=8.5, italic=True, box_w=3.8)

def panel3(f):
    A = Ax(f, [0.14, 0.36, 0.82, 0.58], (0, 1.2), (0, 1.25)); A.spines("d*", "a*")
    A.poly([0, 1], [1.0, 0.0], BLUE, 2.4)
    for x, nm in ((0.4, "d_{SD}"), (0.6, "d_{NC}")):
        y = 1 - x
        A.line((0, 0), (x, y), ORANGE, 2.4)
        A.line((x, 0), (x, y), INK, 1, "dot"); A.line((0, y), (x, y), INK, 1, "dot")
        A.dot((x, y), 0.035, ORANGE)
        A.text((x, 0), nm, dy=0.13, ha="center", size=9.5, italic=True, box_w=0.8)
        A.text((0, y), f"a*({nm})", dx=-0.05, ha="right", size=8.5, italic=True, box_w=0.9)
    A.text((1, 0), "d_0^*", dy=0.13, ha="center", size=9.5, italic=True, box_w=0.8)
    A.dot((1, 0), 0.03, BLUE)
    f.text(0.5, 0.105, "d_{SD} = 0,4 d_0^*     d_{NC} = 0,6 d_0^*", ha="center", size=9, italic=True, box_w=3.8)
    f.text(0.5, 0.03, "T_{LS} = 1,6π √( d_{LS} / a*(d_{LS}) ),  LS = SD, NC", ha="center", size=8.5, italic=True, box_w=3.8)

XM = 0.24; YM = 6.9
def panel4(f):
    A = Ax(f, [0.17, 0.36, 0.79, 0.59], (0, XM), (0, YM)); A.spines("S_d", "S_a")
    A.poly(Sde(TT), Se(TT), NAVY, 2.4)
    T = 1.11; Ses = float(Se(T)); det = Ses * (T / (2 * np.pi)) ** 2
    A.line((0, 0), (det, Ses), INK, 1.2, "dash")
    xc, yc = float(Sde(TC)), float(Se(TC)); A.line((0, 0), (xc, yc), GRAY, 0.9)
    A.dot((det, Ses), 0.035, INK)
    A.line((det, 0), (det, Ses), INK, 1, "dot"); A.line((0, Ses), (det, Ses), INK, 1, "dot")
    A.text((det, 0), "S_De(T_{LS})", dy=0.15, ha="center", size=9, italic=True, box_w=1.1)
    A.text((0, Ses), "S_e(T_{LS})", dx=-0.05, ha="right", size=9, italic=True, box_w=0.9)
    for p, fr, s, c in (((det, Ses), 0.5, "T_{LS}", INK), ((xc, yc), 0.7, "T_C", GRAY)):
        x0, y0 = A.P(0, 0); x1, y1 = A.P(*p); ang = math.degrees(math.atan2(-(y1 - y0), x1 - x0))
        cx, cy = x0 + fr * (x1 - x0), y0 + fr * (y1 - y0)
        nx, ny = math.sin(math.radians(ang)), math.cos(math.radians(ang))
        sg = -1 if s.startswith("T_{") else 1
        S.text(cx + sg * nx * 0.10, cy + sg * ny * 0.10, s, size=10, color=c, ha="center", italic=True, rot=-ang, box_w=0.7)
    A.text((0.17, 6.5), "espectro elástico", color=NAVY, size=8.5, ha="center", box_w=1.4)
    A.text((0.17, 6.0), "(ξ = 5 %)", color=NAVY, size=8.5, ha="center", box_w=1.4)
    f.text(0.5, 0.2, "S_De(T) = S_e(T) (T/2π)^2", ha="center", size=9, italic=True, box_w=3.2)
    f.text(0.5, 0.115, "T_{LS} = 1,6π √( d_{LS} / a*(d_{LS}) )", ha="center", size=9, italic=True, box_w=3.2)
    f.text(0.5, 0.03, "formato aceleración-desplazamiento", ha="center", size=8, box_w=3.2)

def eta(xi): return max(0.55, math.sqrt(10 / (5 + xi)))
def panel5(f):
    a0, d0 = 3.0, 0.22
    cases = [("SD", "ξ = 8 %", 8, 0.4), ("NC", "ξ = 10 %", 10, 0.6)]
    for i, (nm, xt, xi, fr) in enumerate(cases):
        A = Ax(f, [0.135 + i * 0.5, 0.45, 0.34, 0.47], (0, XM), (0, YM)); A.spines("S_d", None)
        e = eta(xi); dls = fr * d0; als = a0 * (1 - fr)
        Tls = 1.6 * math.pi * math.sqrt(dls / als)
        A.poly(Sde(TT), Se(TT), NAVY, 2.2)
        A.poly(e * Sde(TT), e * Se(TT), TEAL, 2.2)
        A.poly([0, d0], [a0, 0], BLUE, 2.4)
        dd = e * float(Sde(Tls)); aa = e * float(Se(Tls))
        A.line((0, 0), (dd, aa), INK, 1, "dash")
        A.dot((dls, als), 0.035, BLUE); A.line((dls, 0), (dls, als), BLUE, 1, "dot")
        A.dot((dd, aa), 0.042, RED); A.line((dd, 0), (dd, aa), RED, 1, "dot")
        ha_d, ha_c = ("right", "left")
        A.text((dd, 0), "Δd", color=RED, dy=0.14, ha=ha_d, size=10, italic=True, bold=True, box_w=0.6)
        A.text((dls, 0), f"d_{{{nm}}}", color=BLUE, dy=0.14, ha=ha_c, size=9.5, italic=True, box_w=0.7)
        A.text((0, YM), f"{nm}   ({xt})", dy=-0.1, ha="left", size=8.5, color=GRAY, italic=True, box_w=1.8)
    LY = 0.285
    for (c, t), x in zip(((NAVY, "Espec. elástico"), (TEAL, "Espec. amort."), (BLUE, "capacidad")), (0.02, 0.255, 0.5)):
        S.line(f.pt(x, LY), f.pt(x + 0.035, LY), c, 3.0); f.text(x + 0.045, LY, t, color=c, size=9, box_w=1.1)
    f.text(0.69, LY, "Amort", color=TEAL, bold=True, size=10, box_w=0.6)
    f.text(0.775, LY, "=", size=10, box_w=0.2)
    f.text(0.805, LY, "Elas", color=NAVY, bold=True, size=10, box_w=0.5)
    f.text(0.905, LY, "· η", bold=True, size=10, box_w=0.5)
    f.text(0.06, 0.205, "Δd = η(ξ) S_De(T_{LS})", size=10, italic=True, box_w=2.6)
    f.text(0.06, 0.135, "η = √( 10 / (5 + ξ) )  ≥ 0,55", size=10, italic=True, box_w=2.8)
    f.text(0.06, 0.065, "ξ_{SD} = 8 %", size=10, italic=True, box_w=1.6)
    f.text(0.25, 0.065, "ξ_{NC} = 10 %", size=10, italic=True, box_w=1.6)

def panel6(f):
    A = EqAx(f, [0.0, 0.2, 0.38, 0.78], (-0.8, 2.6), 0.42, -1.0)
    th = 0.12; c = draw_pier(A, th, loads=False)
    top = 2.9
    A.line((0.5, 2.6), (0.5, top + 0.1), INK, 0.8, "dot"); A.line((c[0], c[1]), (c[0], top + 0.1), INK, 0.8, "dot")
    A.line((0.5, top), (c[0], top), RED, 1.3, arrow="both")
    A.text(((0.5 + c[0]) / 2, top), "d_C", color=RED, bold=True, italic=True, size=11, ha="center", va="bottom", box_w=0.8, dy=-0.02)
    B = Ax(f, [0.52, 0.36, 0.45, 0.55], (0, 1.12), (0, 1.25)); B.spines("d_C", "λ")
    B.poly([0, 1], [0.8, 0.0], GRAY, 2.2)
    ud, ul = 0.33, 0.4
    for x, lab, c_, ha in ((ud, "Γ Δd", RED, "right"), (ul, "Γ d_{LS}", INK, "left")):
        B.line((x, 0), (x, 0.8 * (1 - x)), c_, 1, "dot"); B.dot((x, 0.8 * (1 - x)), 0.034, c_)
        B.text((x, 0), lab, dy=0.13, ha=ha, size=9, italic=True, color=c_, box_w=0.9)
    f.text(0.5, 0.105, "d_C = Γ · Δd", ha="center", size=10.5, italic=True, box_w=3.0)
    f.text(0.5, 0.03, "γ_{Sd} Δd ≤ d_{LS} / γ_{Rd}", ha="center", size=9, italic=True, box_w=3.2)

TAGS = [("≠", "carga uniforme sobre bloque rígido; sin tramo elástico"),
        ("=", "misma transformación (Γ, m*, d* = d_C/Γ); cambia a* = λ g/e*"),
        ("≠", "sin bilineal: oscilador lineal con rigidez secante"),
        ("=", "mismo espectro; se evalúa en T_LS, no en T*"),
        ("≠", "se reduce con η (amortiguamiento), no con R; sin corte"),
        ("≠", "se compara Δd con d_LS; no se exige 1,5 d_t")]
titles = ["Curva de capacidad: cabeceo", "Sistema de un grado de libertad", "Sistema lineal equivalente",
          "Espectro y periodo T_LS", "Desplazamiento de demanda", "De vuelta a la estructura"]
panels = [panel1, panel2, panel3, panel4, panel5, panel6]

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
X0, Y0, GAP = 0.45, 1.0, 0.13
CHS = [2.9, 3.1]
WIDTHS = [[4.011, 4.011, 4.011], [3.45, 5.19, 3.45]]
IMG = [[(3.85, 2.25)] * 3, [(3.25, 2.5), (5.1, 2.5), (3.25, 2.5)]]
S.text(0.45, 0.28, "Método de cabeceo, por analogía con Eurocódigo 8", size=20, bold=True, box_w=12.4, va="top")
for i, (t, fn) in enumerate(zip(titles, panels)):
    r, c = divmod(i, 3)
    CW = WIDTHS[r][c]; x = X0 + sum(WIDTHS[r][:c]) + c * GAP; y = Y0 + (0 if r == 0 else CHS[0] + GAP); CH = CHS[r]
    S.cur = f"Panel {i + 1}"
    S.rect(x, y, CW, CH, BG, radius=0.04)
    S.text(x + 0.16, y + 0.24, str(i + 1), size=18, bold=True, box_w=0.4)
    S.text(x + 0.5, y + 0.22, t, size=13, bold=True, box_w=CW - 0.6, italic=False)
    mark, txt = TAGS[i]
    S.text(x + 0.5, y + 0.45, f"{mark} N2: {txt}", size=9, color=(PLUM if mark == "≠" else EQ), box_w=CW - 0.6)
    w, h = IMG[r][c]
    fn(Fig(x + (CW - w) / 2, y + 0.58, w, h))
S.cur = None
S.text(0.45, 7.3, "EN 1998-3:2025 (11.3.3, 11.4.2, 11.5.2.2) por analogía (excluye puentes de fábrica, 12.1(3)); espectro de UNE-EN 1998-1:2018. DL no se muestra.", size=10, color=MUTED, box_w=12.4)

if __name__ == "__main__":
    render_pptx(slide, S)
    prs.save("Cabeceo_una_diapositiva.pptx")
    render_preview(S, path="preview_cab.png")
    print("ok", len(S.items))
