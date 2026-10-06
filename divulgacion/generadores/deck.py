import math
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from scene import Scene, clip_poly, render_pptx, render_preview

GRAY, BLUE, ORANGE, TEAL, RED = "#8A8F98", "#1F6FB5", "#D9822B", "#17998A", "#C0392B"
INK, NAVY, BG, WHITE, MUTED = "#1F2A44", "#0E2A66", "#F3F5F8", "#FFFFFF", "#5A6373"
def blend(c1, c2, t):
    a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]; b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
    return "#%02X%02X%02X" % tuple(int(round(x * t + y * (1 - t))) for x, y in zip(a, b))
FAINT = blend(INK, BG, 0.55)

S = Scene()

class Fig:
    """Zona de dibujo de un panel: origen (esquina superior izquierda) y tamano en pulgadas."""
    def __init__(self, ox, oy, W, H): self.ox, self.oy, self.W, self.H = ox, oy, W, H
    def pt(self, fx, fy): return (self.ox + fx * self.W, self.oy + (1 - fy) * self.H)   # fracciones, origen abajo-izq
    def text(self, fx, fy, s, **kw): S.text(*self.pt(fx, fy), s, **kw)

class Ax:
    def __init__(self, fig, box, xlim, ylim):
        self.f, self.box, self.xl, self.yl = fig, box, xlim, ylim
    def P(self, x, y):
        l, b, w, h = self.box
        fx = l + (x - self.xl[0]) / (self.xl[1] - self.xl[0]) * w
        fy = b + (y - self.yl[0]) / (self.yl[1] - self.yl[0]) * h
        return self.f.pt(fx, fy)
    def poly(self, xs, ys, color, w=2.0, dash=None, clip=True):
        pts = list(zip(xs, ys))
        segs = clip_poly(pts, self.xl[1], self.yl[1]) if clip else [pts]
        for sg in segs: S.poly([self.P(*p) for p in sg], color, w, dash)
    def line(self, p0, p1, color, w=1.0, dash=None, **kw): S.line(self.P(*p0), self.P(*p1), color, w, dash, **kw)
    def dot(self, p, r, color, **kw): S.dot(self.P(*p), r, color, **kw)
    def text(self, p, s, dx=0, dy=0, **kw):
        x, y = self.P(*p); S.text(x + dx, y + dy, s, **kw)
    def spines(self, xlabel=None, ylabel=None, w=1.0):
        x0, x1 = self.xl; y0, y1 = self.yl
        self.line((x0, y0), (x0, y1), INK, w); self.line((x0, y0), (x1, y0), INK, w)
        if xlabel: self.text((x1, y0), xlabel, dy=0.13, ha="right", size=9, italic=True, box_w=1.2)
        if ylabel: self.text((x0, y1), ylabel, dy=-0.1, ha="center", size=9, italic=True, box_w=1.2)

# ---------------- curva sintetica ----------------
AA, UU = 3.0, 1.6
def curve(u):
    u = np.asarray(u, float)
    up = (1 - np.exp(-AA * np.minimum(u, 1))) / (1 - np.exp(-AA))
    dn = 1 - 0.2 * np.maximum(u - 1, 0) / (UU - 1)
    return np.where(u <= 1, up, dn)

def building(ax, disp=0.0, arrows=True):
    xs = lambda h: h * disp / 3.0
    if disp:
        for x0 in (0.0, 1.0): ax.line((x0, 0), (x0, 3), FAINT, 0.8)
        for h in (1, 2, 3): ax.line((0, h), (1, h), FAINT, 0.8)
        ax.dot((0.5, 3), 0.035, INK, hollow=True, bg=BG)
    for x0 in (0.0, 1.0): ax.poly([x0 + xs(h) for h in range(4)], list(range(4)), INK, 2.0, clip=False)
    for h in (1, 2, 3): ax.line((xs(h), h), (1 + xs(h), h), INK, 2.5)
    ax.line((-0.25, 0), (1.25, 0), INK, 1.5)
    if arrows:
        for h in (1, 2, 3):
            L = 0.2 * h; ax.line((1 + xs(h) + 0.05, h), (1 + xs(h) + 0.05 + L + 0.05, h), RED, 1.5, arrow="end")
    ax.dot((xs(3) + 0.5, 3), 0.042, BLUE)

SPEC = dict(AGS=2.5, TB=0.15, TC=0.5, TD=2.0)
def Se(T):
    T = np.asarray(T, float); a = SPEC["AGS"] * 2.5; TB, TC, TD = SPEC["TB"], SPEC["TC"], SPEC["TD"]
    return np.where(T <= TB, SPEC["AGS"] * (1 + T / TB * 1.5), np.where(T <= TC, a, np.where(T <= TD, a * TC / T, a * TC * TD / T ** 2)))
def Sde(T): return Se(T) * (np.asarray(T, float) / (2 * np.pi)) ** 2
def Rmu(T, mu): return np.where(np.asarray(T) < SPEC["TC"], (mu - 1) * np.asarray(T) / SPEC["TC"] + 1, mu)
TT = np.linspace(0.01, 3.0, 400)
TC = SPEC["TC"]

def numtxt(): pass

# =============== paneles ===============
def panel1(f):
    A = Ax(f, [0.0, 0.2, 0.42, 0.78], (-0.35, 2.3), (-1.2, 4.2)); building(A, 0.55)
    A.line((0.5, 3), (0.5, 3.5), INK, 0.8, "dot"); A.line((1.05, 3), (1.05, 3.5), INK, 0.8, "dot")
    A.line((0.5, 3.4), (1.05, 3.4), BLUE, 1.3, arrow="both")
    A.text((0.78, 3.55), "d_C", color=BLUE, bold=True, italic=True, size=11, ha="center", va="bottom", box_w=0.8)
    A.line((1.35, -0.62), (-0.15, -0.62), RED, 3.4, arrow="end", big=True)
    A.text((0.6, -1.0), "F_b", color=RED, bold=True, italic=True, size=12, ha="center", box_w=0.8)
    B = Ax(f, [0.56, 0.36, 0.42, 0.58], (0, 2.15), (0, 1.25)); B.spines("d_C", "F_b")
    u = np.linspace(0, 2.0, 120); B.poly(u, curve(u), GRAY, 2.4)
    B.line((0, 1), (1, 1), INK, 1, "dot"); B.line((1, 0), (1, 1), INK, 1, "dot"); B.dot((1, 1), 0.03, INK)
    B.line((UU, 0), (UU, 0.8), INK, 1, "dot"); B.dot((UU, 0.8), 0.03, INK)
    B.text((0, 1), "F_{b,máx}", dx=-0.05, ha="right", size=9, italic=True, box_w=1.2)
    B.text((1, 0), "d_m", dy=0.12, ha="center", size=9.5, italic=True, box_w=0.8)
    B.text((UU, 0), "d_u", dy=0.12, ha="center", size=9.5, italic=True, box_w=0.8)
    B.text((1.62, 0.9), "80 %", size=8.5, box_w=0.8)
    f.text(0.5, 0.05, "F_i ∝ m_i Φ_i   (carga proporcional al modo)", ha="center", size=9, box_w=3.6)

G = 1.4
def panel2(f):
    A = Ax(f, [0.14, 0.36, 0.82, 0.58], (0, 2.15), (0, 1.25)); A.spines("d_C, d*", "F_b, F*")
    u = np.linspace(0, 2.0, 120)
    A.poly(u, curve(u), GRAY, 2.0); A.poly(u / G, curve(u) / G, BLUE, 2.6)
    A.line((1 / G, 0), (1 / G, 1 / G), BLUE, 1, "dot"); A.line((0, 1 / G), (1 / G, 1 / G), BLUE, 1, "dot")
    A.dot((1 / G, 1 / G), 0.03, BLUE); A.dot((UU / G, 0.8 / G), 0.03, BLUE); A.line((UU / G, 0), (UU / G, 0.8 / G), BLUE, 1, "dot")
    A.text((0, 1 / G), "F_y^*", color=BLUE, dx=-0.05, ha="right", size=9.5, italic=True, box_w=0.8)
    A.text((1 / G, 0), "d_m^*", color=BLUE, dy=0.13, ha="center", size=9.5, italic=True, box_w=0.8)
    A.text((UU / G, 0), "d_u^*", color=BLUE, dy=0.13, ha="center", size=9.5, italic=True, box_w=0.8)
    A.line((0.75, 0.935), (0.57, 0.75), INK, 1, arrow="end")
    f.text(0.5, 0.105, "F* = F_b/Γ     d* = d_C/Γ     m* = Σ m_i Φ_i", ha="center", size=9, italic=True, box_w=3.8)
    f.text(0.5, 0.03, "Γ = m* / Σ m_i Φ_i^2     (la rigidez no cambia)", ha="center", size=8.5, box_w=3.8)

def panel3(f):
    A = Ax(f, [0.14, 0.36, 0.82, 0.58], (0, 2.15), (0, 1.25)); A.spines("d*", "F*")
    Fy = dm = 1 / G
    ds = np.linspace(0, dm, 80); Fs = curve(ds * G) / G
    Em = np.trapezoid(Fs, ds); dy = 2 * (dm - Em / Fy)
    # areas
    S.polygon([A.P(0, 0)] + [A.P(x, y) for x, y in zip(ds, Fs)] + [A.P(dm, 0)], fill=BLUE, alpha=0.28)
    S.polygon([A.P(0, 0), A.P(dy, Fy), A.P(dm, Fy), A.P(dm, 0)], hatch="wdUpDiag", hatch_color=ORANGE)
    full = np.linspace(0, 2.0 / G, 120); A.poly(full, curve(full * G) / G, BLUE, 2.0)
    A.poly([0, dy, UU / G], [0, Fy, Fy], ORANGE, 2.6)
    A.line((dm, 0), (dm, Fy), INK, 1, "dot"); A.line((dy, 0), (dy, Fy), INK, 1, "dot"); A.line((0, Fy), (dm, Fy), INK, 1, "dot")
    A.text((0, Fy), "F_y^*", dx=-0.05, ha="right", size=9.5, italic=True, box_w=0.8)
    A.text((dy, 0), "d_y^*", color=ORANGE, dy=0.13, ha="center", size=9.5, italic=True, box_w=0.8)
    A.text((dm, 0), "d_m^*", dy=0.13, ha="center", size=9.5, italic=True, box_w=0.8)
    A.text((dy * 0.12, Fy * 0.75), "k*", color=ORANGE, size=10, italic=True, box_w=0.6)
    A.text((dm * 0.55, Fy * 1.1), "áreas iguales", size=8.5, ha="center", box_w=1.4)
    f.text(0.5, 0.105, "F_y^* = F*(d_m^*)     d_y^* = 2 (d_m^* − E_m^*/F_y^*)", ha="center", size=9, italic=True, box_w=3.8)
    f.text(0.5, 0.03, "E_m^*: área bajo la curva real hasta d_m^*", ha="center", size=8.5, box_w=3.8)

def along(A, p, frac, s, color, off=0.09, size=10):
    x0, y0 = A.P(0, 0); x1, y1 = A.P(*p)
    ang = math.degrees(math.atan2(-(y1 - y0), x1 - x0))   # angulo visual (positivo = hacia arriba)
    cx, cy = x0 + frac * (x1 - x0), y0 + frac * (y1 - y0)
    nx, ny = -math.sin(math.radians(ang)), -math.cos(math.radians(ang))   # normal hacia arriba-izquierda
    S.text(cx - nx * 0 + nx * off * -1 * -1, cy + ny * off * -1 * -1, s, size=size, color=color, ha="center", italic=True, rot=-ang, box_w=0.6)

def panel4(f):
    xm, ym = 0.075, 7.6
    A = Ax(f, [0.17, 0.36, 0.79, 0.59], (0, xm), (0, ym)); A.spines("S_d", "S_a")
    A.poly(Sde(TT), Se(TT), NAVY, 2.4)
    Ts, ay = 0.8, 1.6
    Ses = float(Se(Ts)); det = Ses * (Ts / (2 * np.pi)) ** 2; dy = ay * (Ts / (2 * np.pi)) ** 2
    A.poly([0, dy, xm], [0, ay, ay], ORANGE, 2.4)
    A.line((0, 0), (det, Ses), INK, 1.2, "dash")
    xc, yc = float(Sde(TC)), float(Se(TC)); A.line((0, 0), (xc, yc), GRAY, 0.9)
    A.dot((det, Ses), 0.035, INK)
    A.line((det, 0), (det, Ses), INK, 1, "dot"); A.line((0, Ses), (det, Ses), INK, 1, "dot")
    A.text((det, 0), "d_{et}^*", dy=0.15, ha="center", size=9.5, italic=True, box_w=0.8)
    A.text((0, Ses), "S_e(T*)", dx=-0.05, ha="right", size=9, italic=True, box_w=0.9)
    A.text((0, ay), "F_y^*/m*", color=ORANGE, dx=-0.05, ha="right", size=9, italic=True, box_w=0.9)
    # rotulos sobre las rectas
    for p, fr, s, c in (((det, Ses), 0.5, "T*", INK), ((xc, yc), 0.3, "T_C", GRAY)):
        x0, y0 = A.P(0, 0); x1, y1 = A.P(*p); ang = math.degrees(math.atan2(-(y1 - y0), x1 - x0))
        cx, cy = x0 + fr * (x1 - x0), y0 + fr * (y1 - y0)
        nx, ny = math.sin(math.radians(ang)), math.cos(math.radians(ang))
        S.text(cx - nx * 0.10, cy - ny * 0.10, s, size=10, color=c, ha="center", italic=True, rot=-ang, box_w=0.6)
    A.text((0.052, 6.9), "espectro elástico", color=NAVY, size=8.5, ha="center", box_w=1.4)
    A.text((0.052, 6.4), "(ξ = 5 %)", color=NAVY, size=8.5, ha="center", box_w=1.4)
    f.text(0.5, 0.2, "d_{et}^* = S_e(T*) (T*/2π)^2", ha="center", size=9, italic=True, box_w=3.2)
    f.text(0.5, 0.115, "T* = 2π √( d_y^* / (F_y^*/m*) )", ha="center", size=9, italic=True, box_w=3.2)
    f.text(0.5, 0.03, "formato aceleración-desplazamiento (F*/m*)", ha="center", size=8, box_w=3.2)

def panel5(f):
    cases = [("T* ≥ T_C", 0.8, 1.6), ("T* < T_C", 0.25, 2.5)]
    xm, ym = 0.095, 7.2
    for i, (title, Ts, ay) in enumerate(cases):
        A = Ax(f, [0.135 + i * 0.5, 0.45, 0.34, 0.47], (0, xm), (0, ym)); A.spines("S_d", None)
        Ses = float(Se(Ts)); det = Ses * (Ts / (2 * np.pi)) ** 2
        qu = Ses / ay; mu = qu if Ts >= TC else 1 + (qu - 1) * TC / Ts
        dy = ay * (Ts / (2 * np.pi)) ** 2; dt = mu * dy
        R = Rmu(TT, mu)
        A.poly(Sde(TT), Se(TT), NAVY, 2.2)
        A.poly(mu / R * Sde(TT), Se(TT) / R, TEAL, 2.2)
        A.poly([0, dy, xm], [0, ay, ay], ORANGE, 2.4)
        A.line((0, 0), (det, Ses), INK, 1, "dash"); A.dot((det, Ses), 0.032, INK)
        A.line((0, ay), (xm, ay), ORANGE, 0.8, "dot"); A.line((dy, 0), (dy, ay), INK, 0.9, "dot"); A.line((dt, 0), (dt, ay), RED, 1, "dot")
        A.dot((dt, ay), 0.042, RED)
        A.text((0, ay), "F_y^*/m*", color=ORANGE, dx=-0.05, ha="right", size=8.5, italic=True, box_w=0.9)
        ha_y, ha_t = ("center", "center") if i == 0 else ("right", "left")
        A.text((dy, 0), "d_y^*", color=ORANGE, dy=0.14, ha=ha_y, size=9.5, italic=True, box_w=0.7)
        A.text((dt, 0), "d_t^*", color=RED, dy=0.14, ha=ha_t, size=10, italic=True, bold=True, box_w=0.7)
        A.text((0, ym), title, dy=-0.1, dx=0.0, ha="left", size=8.5, color=GRAY, italic=True, box_w=1.5)
    LY = 0.285
    for (c, t), x in zip(((NAVY, "elástico"), (TEAL, "inelástico"), (ORANGE, "bilineal")), (0.06, 0.23, 0.41)):
        S.line(f.pt(x, LY), f.pt(x + 0.035, LY), c, 3.0); f.text(x + 0.045, LY, t, color=c, size=9, box_w=1.0)
    f.text(0.62, LY, "Inel", color=TEAL, bold=True, size=10, box_w=0.5)
    f.text(0.675, LY, "=", size=10, box_w=0.2)
    f.text(0.695, LY, "Elas", color=NAVY, bold=True, size=10, box_w=0.5)
    f.text(0.76, LY, "/ R", bold=True, size=10, box_w=0.5)
    f.text(0.06, 0.205, "μ = d_t^* / d_y^*  ≥ 1", size=10, italic=True, box_w=2.6)
    f.text(0.06, 0.135, "R = μ", size=10, italic=True, box_w=2.0)
    f.text(0.60, 0.135, "T ≥ T_C", size=10, italic=True, color=GRAY, ha="right", box_w=1.0)
    f.text(0.06, 0.065, "R = 1 + (μ − 1) T / T_C", size=10, italic=True, box_w=2.6)
    f.text(0.60, 0.065, "T < T_C", size=10, italic=True, color=GRAY, ha="right", box_w=1.0)

def panel6(f):
    A = Ax(f, [0.0, 0.2, 0.36, 0.78], (-0.35, 2.3), (-1.2, 4.2)); building(A, 0.55, arrows=False)
    A.line((0.5, 3), (0.5, 3.4), INK, 0.8, "dot"); A.line((1.05, 3), (1.05, 3.4), INK, 0.8, "dot")
    A.line((0.5, 3.3), (1.05, 3.3), RED, 1.3, arrow="both")
    A.text((0.78, 3.5), "d_t", color=RED, bold=True, italic=True, size=11, ha="center", va="bottom", box_w=0.8)
    B = Ax(f, [0.5, 0.36, 0.48, 0.55], (0, 2.15), (0, 1.25)); B.spines("d_C", "F_b")
    u = np.linspace(0, 2.0, 120); B.poly(u, curve(u), GRAY, 2.2)
    ut = 0.85; um = np.linspace(0, 1.5 * ut, 60)
    S.polygon([B.P(0, 0)] + [B.P(x, float(curve(x))) for x in um] + [B.P(1.5 * ut, 0)], fill=RED, alpha=0.12)
    for x, lab, c in ((ut, "d_t", RED), (1.5 * ut, "1,5 d_t", INK)):
        B.line((x, 0), (x, float(curve(x))), c, 1, "dot"); B.dot((x, float(curve(x))), 0.032, c)
        B.text((x, 0), lab, dy=0.13, ha="center", size=9, italic=True, color=c, box_w=0.9)
    B.line((0.02, 0.14), (1.5 * ut, 0.14), INK, 1, arrow="end")
    f.text(0.5, 0.105, "d_t = Γ d_t^*", ha="center", size=10.5, italic=True, box_w=3.0)
    f.text(0.5, 0.03, "el pushover debe llegar al menos a 1,5 d_t", ha="center", size=8.5, box_w=3.2)

# =============== montaje de la diapositiva ===============
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
X0, Y0, GAP = 0.45, 1.02, 0.15
CHS = [2.91, 3.1]
WIDTHS = [[4.011, 4.011, 4.011], [3.5, 5.03, 3.5]]
IMG = [[(3.85, 2.4)] * 3, [(3.3, 2.6), (5.0, 2.6), (3.3, 2.6)]]
titles = ["Pushover", "Sistema de un grado de libertad", "Bilineal equivalente",
          "Demanda: espectro y periodo T*", "Desplazamiento objetivo", "De vuelta a la estructura"]
panels = [panel1, panel2, panel3, panel4, panel5, panel6]
S.text(0.45, 0.45, "Método N2: de la estructura al desplazamiento objetivo", size=30, bold=True, box_w=12.4)
S.text(0.45, 0.83, "UNE-EN 1998-1:2018, anejo B · esquema ilustrativo, no a escala", size=13, color=MUTED, box_w=12.4)
for i, (t, fn) in enumerate(zip(titles, panels)):
    r, c = divmod(i, 3)
    CW = WIDTHS[r][c]; x = X0 + sum(WIDTHS[r][:c]) + c * GAP; y = Y0 + (0 if r == 0 else CHS[0] + GAP); CH = CHS[r]
    S.cur = f"Panel {i + 1}"
    S.rect(x, y, CW, CH, BG, radius=0.04)
    S.text(x + 0.16, y + 0.26, str(i + 1), size=18, bold=True, box_w=0.4)
    S.text(x + 0.5, y + 0.26, t, size=13, bold=True, box_w=CW - 0.6)
    w, h = IMG[r][c]
    fn(Fig(x + (CW - w) / 2, y + 0.46, w, h))
S.cur = None
S.text(0.45, 7.31, "Nota, iteración opcional (anejo B): si d_t^* difiere mucho de d_m^*, se repiten los pasos 3 a 5 con d_m^* = d_t^* hasta converger.", size=12, color=MUTED, box_w=12.4)

if __name__ == "__main__":
    render_pptx(slide, S)
    prs.save("N2_editable.pptx")
    render_preview(S, path="preview.png")
    print("ok", len(S.items))
