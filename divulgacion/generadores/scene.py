"""Primitivas de dibujo con dos salidas: formas nativas de PowerPoint (editables) y una vista previa PNG (PIL)."""
import math, re
import numpy as np
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
def _hex(c): return c.lstrip("#").upper()

# ---------------- texto con marcas: _x subindice, ^x superindice, _{..} ^{..} ----------------
def parse_rt(s):
    """'F_{y}^{*}/m^{*}' -> [('F','n'),('y','sub'),('*','sup'),('/m','n'),('*','sup')]"""
    out, i, buf = [], 0, ""
    def flush():
        nonlocal buf
        if buf: out.append((buf, "n")); buf = ""
    while i < len(s):
        ch = s[i]
        if ch in "_^":
            flush(); mode = "sub" if ch == "_" else "sup"; i += 1
            if i < len(s) and s[i] == "{":
                j = s.index("}", i); tok = s[i + 1:j]; i = j + 1
            else:
                tok = s[i]; i += 1
            out.append((tok, mode))
        else:
            buf += ch; i += 1
    flush(); return out

class Scene:
    def __init__(self): self.items = []; self.cur = None
    def add(self, kind, **kw): kw["group"] = self.cur; self.items.append((kind, kw))
    def poly(self, pts, color, w=2.0, dash=None, group=None): self.add("poly", pts=pts, color=color, w=w, dash=dash, group=group)
    def polygon(self, pts, fill=None, alpha=None, hatch=None, hatch_color=None, group=None):
        self.add("polygon", pts=pts, fill=fill, alpha=alpha, hatch=hatch, hatch_color=hatch_color, group=group)
    def line(self, p0, p1, color, w=1.0, dash=None, arrow=None, big=False, group=None):
        self.add("line", p0=p0, p1=p1, color=color, w=w, dash=dash, arrow=arrow, big=big, group=group)
    def dot(self, c, r, color, hollow=False, bg=None, group=None): self.add("dot", c=c, r=r, color=color, hollow=hollow, bg=bg, group=group)
    def text(self, x, y, s, size=9, color="#1F2A44", ha="left", va="middle", bold=False, italic=False, rot=0, box_w=2.6, group=None):
        self.add("text", x=x, y=y, s=s, size=size, color=color, ha=ha, va=va, bold=bold, italic=italic, rot=rot, box_w=box_w, group=group)
    def rect(self, x, y, w, h, fill, radius=None, group=None): self.add("rect", x=x, y=y, w=w, h=h, fill=fill, radius=radius, group=group)

def clip_poly(pts, xmax=None, ymax=None):
    """Recorta una poligonal a x<=xmax, y<=ymax (en coordenadas de datos). Devuelve lista de tramos."""
    def inside(p): return (xmax is None or p[0] <= xmax + 1e-12) and (ymax is None or p[1] <= ymax + 1e-12)
    def cross(p, q):
        t = 1.0
        if xmax is not None and (p[0] - xmax) * (q[0] - xmax) < 0: t = min(t, (xmax - p[0]) / (q[0] - p[0]))
        if ymax is not None and (p[1] - ymax) * (q[1] - ymax) < 0: t = min(t, (ymax - p[1]) / (q[1] - p[1]))
        return (p[0] + t * (q[0] - p[0]), p[1] + t * (q[1] - p[1]))
    out, cur = [], []
    for i, p in enumerate(pts):
        if inside(p):
            if not cur and i > 0 and not inside(pts[i - 1]): cur.append(cross(p, pts[i - 1]))
            cur.append(p)
        else:
            if cur: cur.append(cross(cur[-1], p)); out.append(cur); cur = []
    if cur: out.append(cur)
    return out

# =========================== PowerPoint ===========================
DASH = {"dot": MSO_LINE.ROUND_DOT, "dash": MSO_LINE.DASH, None: None}
def _style_line(line, color, w, dash):
    line.color.rgb = RGBColor.from_string(_hex(color)); line.width = Pt(w)
    if dash: line.dash_style = DASH[dash]
def _arrow(line, arrow, big):
    ln = line._get_or_add_ln()
    sz = "lg" if big else "med"
    if arrow in ("start", "both"):
        e = etree.SubElement(ln, qn("a:headEnd")); e.set("type", "triangle"); e.set("w", sz); e.set("len", sz)
    if arrow in ("end", "both"):
        e = etree.SubElement(ln, qn("a:tailEnd")); e.set("type", "triangle"); e.set("w", sz); e.set("len", sz)
def _alpha_fill(shape, color, alpha):
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(_hex(color))
    if alpha is not None:
        clr = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        a = etree.SubElement(clr, qn("a:alpha")); a.set("val", str(int(alpha * 100000)))
def _hatch_fill(shape, color, prst="wdUpDiag"):
    spPr = shape._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:pattFill", "a:gradFill"):
        for e in spPr.findall(qn(tag)): spPr.remove(e)
    patt = etree.Element(qn("a:pattFill")); patt.set("prst", prst)
    fg = etree.SubElement(patt, qn("a:fgClr")); c = etree.SubElement(fg, qn("a:srgbClr")); c.set("val", _hex(color))
    bg = etree.SubElement(patt, qn("a:bgClr")); c2 = etree.SubElement(bg, qn("a:srgbClr")); c2.set("val", "FFFFFF")
    a = etree.SubElement(c2, qn("a:alpha")); a.set("val", "0")
    # el relleno va justo despues de la geometria
    geom = spPr.find(qn("a:custGeom"))
    if geom is None: geom = spPr.find(qn("a:prstGeom"))
    geom.addnext(patt)

def _runs(par, parts, size, color, bold, italic):
    for t, mode in parts:
        r = par.add_run(); r.text = t; f = r.font
        f.size = Pt(size); f.bold = bold; f.name = "Calibri"; f.color.rgb = RGBColor.from_string(_hex(color))
        f.italic = italic
        if mode == "sub": f._element.set("baseline", "-25000")
        if mode == "sup": f._element.set("baseline", "30000")

def render_pptx(slide, scene, offset=(0, 0)):
    ox, oy = offset
    E = lambda v: Emu(int(round(v * 914400)))
    shapes = slide.shapes
    groups = {}
    for kind, k in scene.items:
        n0 = len(shapes)
        if kind == "rect":
            sh = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if k["radius"] else MSO_SHAPE.RECTANGLE, E(k["x"]), E(k["y"]), E(k["w"]), E(k["h"]))
            if k["radius"]: sh.adjustments[0] = k["radius"]
            sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(_hex(k["fill"])); sh.line.fill.background(); sh.shadow.inherit = False
        elif kind in ("poly", "polygon"):
            pts = [(int(round((x + ox) * 914400)), int(round((y + oy) * 914400))) for x, y in k["pts"]]
            fb = shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
            fb.add_line_segments(pts[1:], close=(kind == "polygon"))
            sh = fb.convert_to_shape()
            sh.shadow.inherit = False
            if kind == "poly":
                sh.fill.background(); _style_line(sh.line, k["color"], k["w"], k["dash"])
                ln = sh.line._get_or_add_ln(); j = etree.SubElement(ln, qn("a:round"))
            else:
                if k["hatch"]: sh.fill.solid(); _hatch_fill(sh, k["hatch_color"], k["hatch"])
                else: _alpha_fill(sh, k["fill"], k["alpha"])
                sh.line.fill.background()
        elif kind == "line":
            c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(k["p0"][0] + ox), E(k["p0"][1] + oy), E(k["p1"][0] + ox), E(k["p1"][1] + oy))
            _style_line(c.line, k["color"], k["w"], k["dash"])
            if k["arrow"]: _arrow(c.line, k["arrow"], k["big"])
        elif kind == "dot":
            cx, cy = k["c"]; r = k["r"]
            sh = shapes.add_shape(MSO_SHAPE.OVAL, E(cx + ox - r), E(cy + oy - r), E(2 * r), E(2 * r)); sh.shadow.inherit = False
            if k["hollow"]:
                sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(_hex(k["bg"])); _style_line(sh.line, k["color"], 1.0, None)
            else:
                sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(_hex(k["color"])); sh.line.fill.background()
        elif kind == "text":
            bw = k["box_w"]; bh = 0.34
            x = k["x"] + ox; y = k["y"] + oy
            left = x if k["ha"] == "left" else (x - bw / 2 if k["ha"] == "center" else x - bw)
            top = {"middle": y - bh / 2, "top": y, "bottom": y - bh}[k["va"]]
            tb = shapes.add_textbox(E(left), E(top), E(bw), E(bh)); tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.word_wrap = False
            tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}[k["va"]]
            p = tf.paragraphs[0]; p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[k["ha"]]
            _runs(p, parse_rt(k["s"]), k["size"], k["color"], k["bold"], k["italic"])
            if k["rot"]: tb.rotation = k["rot"]
        if k["group"] is not None:
            groups.setdefault(k["group"], []).extend(list(shapes)[n0:])
    for name, shs in groups.items():
        g = shapes.add_group_shape(shapes=shs); g.name = name

# =========================== Vista previa (PIL) ===========================
def render_preview(scene, W=13.333, H=7.5, dpi=140, path="preview.png", bg="#FFFFFF"):
    from PIL import Image, ImageDraw, ImageFont
    import matplotlib
    fd = matplotlib.get_data_path() + "/fonts/ttf/"
    def font(sz, bold=False, italic=False):
        name = "DejaVuSans" + ("-BoldOblique" if bold and italic else "-Bold" if bold else "-Oblique" if italic else "")
        return ImageFont.truetype(fd + name + ".ttf", max(6, int(round(sz * dpi / 72 * 0.88))))
    img = Image.new("RGB", (int(W * dpi), int(H * dpi)), bg); d = ImageDraw.Draw(img, "RGBA")
    P = lambda p: (p[0] * dpi, p[1] * dpi)
    rgb = lambda c, a=255: tuple(int(c.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4)) + (a,)
    def dashed(p0, p1, col, wpx, kind):
        L = math.hypot(p1[0] - p0[0], p1[1] - p0[1]);
        if L == 0: return
        on, off = (wpx * 1.2, wpx * 2.2) if kind == "dot" else (wpx * 5, wpx * 3.5)
        t = 0
        while t < L:
            a = t / L; b = min(t + on, L) / L
            d.line([(p0[0] + a * (p1[0] - p0[0]), p0[1] + a * (p1[1] - p0[1])), (p0[0] + b * (p1[0] - p0[0]), p0[1] + b * (p1[1] - p0[1]))], fill=col, width=max(1, int(wpx)))
            t += on + off
    def arrowhead(p0, p1, col, wpx, big):
        ang = math.atan2(p1[1] - p0[1], p1[0] - p0[0]); s = wpx * (4.5 if big else 3.2) + 5
        pts = [p1, (p1[0] - s * math.cos(ang - 0.4), p1[1] - s * math.sin(ang - 0.4)), (p1[0] - s * math.cos(ang + 0.4), p1[1] - s * math.sin(ang + 0.4))]
        d.polygon(pts, fill=col)
    for kind, k in scene.items:
        if kind == "rect":
            box = [k["x"] * dpi, k["y"] * dpi, (k["x"] + k["w"]) * dpi, (k["y"] + k["h"]) * dpi]
            if k["radius"]: d.rounded_rectangle(box, radius=k["radius"] * min(k["w"], k["h"]) * dpi, fill=rgb(k["fill"]))
            else: d.rectangle(box, fill=rgb(k["fill"]))
        elif kind == "poly":
            col = rgb(k["color"]); wpx = k["w"] * dpi / 72
            pts = [P(p) for p in k["pts"]]
            for a, b in zip(pts[:-1], pts[1:]):
                if k["dash"]: dashed(a, b, col, wpx, k["dash"])
                else: d.line([a, b], fill=col, width=max(1, int(round(wpx))))
            if not k["dash"]:
                for p in pts[1:-1]: d.ellipse([p[0] - wpx / 2, p[1] - wpx / 2, p[0] + wpx / 2, p[1] + wpx / 2], fill=col)
        elif kind == "polygon":
            pts = [P(p) for p in k["pts"]]
            if k["hatch"]:
                m = Image.new("L", img.size, 0); ImageDraw.Draw(m).polygon(pts, fill=255)
                h = Image.new("RGBA", img.size, (0, 0, 0, 0)); hd = ImageDraw.Draw(h)
                for off in range(-img.size[1], img.size[0], 9): hd.line([(off, img.size[1]), (off + img.size[1], 0)], fill=rgb(k["hatch_color"]), width=2)
                img.paste(h, (0, 0), Image.composite(h.split()[3], Image.new("L", img.size, 0), m))
                d = ImageDraw.Draw(img, "RGBA")
            else:
                d.polygon(pts, fill=rgb(k["fill"], int((k["alpha"] if k["alpha"] is not None else 1) * 255)))
        elif kind == "line":
            col = rgb(k["color"]); wpx = k["w"] * dpi / 72; a, b = P(k["p0"]), P(k["p1"])
            if k["dash"]: dashed(a, b, col, wpx, k["dash"])
            else: d.line([a, b], fill=col, width=max(1, int(round(wpx))))
            if k["arrow"] in ("end", "both"): arrowhead(a, b, col, wpx, k["big"])
            if k["arrow"] in ("start", "both"): arrowhead(b, a, col, wpx, k["big"])
        elif kind == "dot":
            c = P(k["c"]); r = k["r"] * dpi
            if k["hollow"]: d.ellipse([c[0] - r, c[1] - r, c[0] + r, c[1] + r], fill=rgb(k["bg"]), outline=rgb(k["color"]), width=2)
            else: d.ellipse([c[0] - r, c[1] - r, c[0] + r, c[1] + r], fill=rgb(k["color"]))
        elif kind == "text":
            parts = parse_rt(k["s"]); sz = k["size"]
            fs = {"n": font(sz, k["bold"], k["italic"]), "sub": font(sz * 0.7, k["bold"], k["italic"]), "sup": font(sz * 0.7, k["bold"], k["italic"])}
            wtot = sum(d.textlength(t, font=fs[m]) for t, m in parts)
            h = font(sz).size * 1.5
            tmp = Image.new("RGBA", (int(wtot + 20), int(h * 2)), (0, 0, 0, 0)); td = ImageDraw.Draw(tmp)
            x = 10; base = h * 0.5
            for t, m in parts:
                dy = {"n": 0, "sub": h * 0.18, "sup": -h * 0.3}[m]
                td.text((x, base + dy), t, font=fs[m], fill=rgb(k["color"])); x += td.textlength(t, font=fs[m])
            ax = {"left": 10, "center": 10 + wtot / 2, "right": 10 + wtot}[k["ha"]]
            ay = {"middle": base + font(sz).size * 0.6, "top": base, "bottom": base + font(sz).size * 1.2}[k["va"]]
            if k["rot"]:
                tmp = tmp.rotate(-k["rot"], center=(ax, ay), resample=Image.BICUBIC)
            img.paste(tmp, (int(k["x"] * dpi - ax), int(k["y"] * dpi - ay)), tmp)
            d = ImageDraw.Draw(img, "RGBA")
    img.save(path)
